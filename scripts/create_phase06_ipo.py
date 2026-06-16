"""Phase 06 — Scale & IPO deliverables.

Generates:
  - IPO-Readiness-Playbook.docx
  - Subsidiary-Spinoff-Playbook.docx
  - Sukuk-and-Green-Bond-Issuance-Guide.docx
  - Carve-Out-and-Divestiture-Playbook.docx
"""
from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path(__file__).parent.parent / "06_Scale-IPO"
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GOLD = RGBColor(0xC8, 0xA0, 0x4A)
GREY = RGBColor(0x55, 0x55, 0x55)


def new_doc():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.9); s.right_margin = Inches(0.9)
    return doc

def H(doc, text, level=1, color=NAVY, size=None):
    p = doc.add_heading("", level=level)
    r = p.add_run(text); r.bold = True; r.font.color.rgb = color
    if size: r.font.size = Pt(size)

def P(doc, text, bold=False, italic=False, color=None, size=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    if align is not None: p.alignment = align

def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")

def cover(doc, title, subtitle):
    P(doc, ""); P(doc, "")
    P(doc, "AFANO GROUP HOLDINGS", bold=True, color=NAVY, size=28, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, title, bold=True, color=NAVY, size=22, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, subtitle, italic=True, color=GOLD, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, f"Effective: {date.today().isoformat()}  ·  Version 1.0",
      color=GREY, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def ipo_readiness():
    doc = new_doc()
    cover(doc, "IPO Readiness Playbook",
          "NYSE · NASDAQ · LSE · NGX · JSE · SGX · Tadawul · Dual Listing")
    H(doc, "1. Strategic Question", 1)
    P(doc, "IPO is one of three exits. The Group may also (a) hold "
           "indefinitely and run as Berkshire Hathaway does, (b) spin off "
           "to existing shareholders (Section 355 tax-free if US), or "
           "(c) sell to strategic or financial buyer. The decision is "
           "subsidiary-by-subsidiary, based on valuation arbitrage, "
           "capital needs, and brand-independence goals.")
    H(doc, "2. 24-Month IPO Timeline", 1)
    months = [
        ("T-24 to T-18", "Audit upgrade (Big Four), governance refresh, "
                          "minutes review, related-party transaction cleanup, "
                          "ESG/SASB baseline."),
        ("T-18 to T-12", "Banker selection (lead, joint-lead, co-managers), "
                          "kickoff meeting, valuation testing, "
                          "S-1/F-1 drafting begins."),
        ("T-12 to T-6",  "First S-1 filing (confidential JOBS Act), SEC "
                          "comments, financial print, testing-the-waters "
                          "meetings with institutional investors."),
        ("T-6 to T-3",   "Public filing, roadshow planning, pricing range, "
                          "syndicate book-building."),
        ("T-3 to T-0",   "Roadshow, pricing, allocation, IPO trade, "
                          "stabilization (greenshoe / overallotment)."),
        ("T+0 to T+6",   "First earnings call, lock-up expiry planning, "
                          "post-IPO IR program established."),
    ]
    for m, b in months:
        bullet(doc, f"{m}: {b}")
    H(doc, "3. Listing Venue Selection", 1)
    bullet(doc, "NYSE — large-cap, global brand visibility, deep institutional ownership. Best for >$1B market cap.")
    bullet(doc, "NASDAQ — technology-tilted, more permissive on dual-class voting structures. Best for SBG 3 spinoffs.")
    bullet(doc, "London Stock Exchange — preferred for mining, resources, specialty insurance, international listings.")
    bullet(doc, "Nigerian Exchange (NGX) — domestic capital, supportive of local-content positioning, smaller free-float.")
    bullet(doc, "Johannesburg Stock Exchange (JSE) — pan-African capital, secondary listing partner for NGX or LSE.")
    bullet(doc, "Singapore Exchange (SGX) — REITs, business trusts, APAC institutional capital.")
    bullet(doc, "Tadawul (Saudi) — Saudi-anchored subsidiaries; sovereign and regional capital.")
    H(doc, "4. Dual-Listing Strategy", 1)
    P(doc, "Primary listing in mature market for liquidity (NYSE / LSE) plus "
           "secondary listing in domestic market for local-investor access "
           "and political-stakeholder alignment. Use Depository Receipt "
           "programs (ADRs, GDRs) where direct listing is impractical.")
    H(doc, "5. Post-IPO Discipline", 1)
    bullet(doc, "Quarterly earnings, annual report, Form 20-F or 10-K depending on filer status.")
    bullet(doc, "SOX 404 internal controls attestation.")
    bullet(doc, "Disclosure controls and Regulation FD training for executives and IR.")
    bullet(doc, "Continuous-disclosure obligation under local listing rules.")
    bullet(doc, "Insider-trading window calendar; pre-clearance for affiliates.")
    out = BASE / "IPO-Readiness-Playbook.docx"
    doc.save(out); print(f"WROTE {out}")


def spinoff():
    doc = new_doc()
    cover(doc, "Subsidiary Spinoff Playbook",
          "Tax-Free Distribution · Carve-Out IPO · Reverse Morris Trust")
    sections = [
        ("1. Spinoff Mechanics",
         "A spinoff distributes shares of a subsidiary to AFANO shareholders "
         "pro-rata. If structured under IRC Section 355, the distribution is "
         "tax-free to both the parent and shareholders, subject to the "
         "five-year active-trade-or-business test and absence of a "
         "disqualifying acquisition within 2 years."),
        ("2. Strategic Rationale",
         "Conglomerate discount removal, focus, separate currency for M&A, "
         "independent capital structure, talent attraction through "
         "subsidiary-specific equity. Trade-offs: loss of cross-subsidy, "
         "stranded overhead, higher consolidated cost of capital."),
        ("3. Pre-Spinoff Operational Separation",
         "12-24 month workstream covering IT (separate ERP, identity, "
         "endpoints), real estate (lease assignment or new leases), "
         "personnel (transfer or new hire), procurement (separate master "
         "agreements), brand (independent already under House of Brands), "
         "tax (transfer-pricing exit, treaty re-routing)."),
        ("4. Transition Services Agreements (TSAs)",
         "Group provides services to SpinCo for 12-24 months on cost-plus "
         "basis. TSA workstreams per function with monthly tracker; aim "
         "for full separation by Month 18."),
        ("5. Spinoff Documentation",
         "Separation Agreement, TSA, IP Assignment Agreement, Tax Matters "
         "Agreement, Employee Matters Agreement, Commercial Agreements "
         "(supply, license, lease). Filed as Form 10 / equivalents."),
        ("6. Reverse Morris Trust Variants",
         "Where SpinCo is to merge with a third party post-spinoff in a "
         "tax-efficient manner, structure as Reverse Morris Trust: AFANO "
         "shareholders end up majority of the combined entity."),
        ("7. Brand and Endorsement After Spinoff",
         "SpinCo retains independent brand. Endorsement footer is removed "
         "from SpinCo's properties as of the spinoff effective date. "
         "Joint press release announces independent listing."),
    ]
    for t, b in sections:
        H(doc, t, 1); P(doc, b)
    out = BASE / "Subsidiary-Spinoff-Playbook.docx"
    doc.save(out); print(f"WROTE {out}")


def sukuk_green():
    doc = new_doc()
    cover(doc, "Sukuk and Green Bond Issuance Guide",
          "Shariah-Compliant Capital · Climate-Aligned Debt · Sustainable Finance")
    sections = [
        ("1. Sukuk — Strategic Rationale",
         "Access to OIC sovereign-investor and Islamic-finance capital "
         "pools. Particularly relevant for SBGs 2 (Energy), 4 (Infrastructure), "
         "and 12 (Industrial). Sukuk Ijara, Wakala, and Mudaraba are the "
         "preferred structures depending on underlying asset and use of "
         "proceeds."),
        ("2. Sukuk Structuring",
         "Use a SPV (Cayman, Labuan, or DIFC-domiciled) as Issuer. "
         "Underlying asset transferred to SPV under sale-and-leaseback or "
         "service-agency arrangement. Returns to investors are profit-share "
         "or lease rentals, not interest. Shariah Supervisory Board sign-off "
         "from a reputable firm (e.g. Dar al Sharia, AAOIFI-certified "
         "scholars)."),
        ("3. Sukuk Distribution Venues",
         "Tadawul (Saudi), Nasdaq Dubai, Bursa Malaysia, LSE. Primary "
         "investors: GCC sovereign and quasi-sovereign, Malaysian "
         "Islamic banks, Turkish katilim banks, Indonesian Islamic "
         "funds."),
        ("4. Green Bond Framework",
         "Aligned with ICMA Green Bond Principles. Use of proceeds "
         "categories: renewable energy, energy efficiency, clean "
         "transportation, sustainable water management, climate change "
         "adaptation, green buildings, sustainable agriculture."),
        ("5. Sustainability-Linked Bond (SLB)",
         "Coupon step-up tied to KPIs (e.g. Scope 1+2 emissions intensity, "
         "renewable energy share, water withdrawal reduction). Less "
         "restrictive than green bond — proceeds may fund general purposes."),
        ("6. External Verification",
         "Pre-issuance Second Party Opinion (Sustainalytics, ISS, Moody's "
         "ESG, Vigeo Eiris). Post-issuance impact reporting annually. "
         "Climate Bonds Initiative certification optional but increasingly "
         "expected by Tier-1 investors."),
        ("7. Issuance Calendar",
         "Aim for 1 benchmark issuance per 18-month cycle. Build the "
         "credit curve gradually: 5-year, then 10-year, then 30-year. "
         "Maintain SEC Rule 144A and Reg S documentation for global "
         "distribution."),
    ]
    for t, b in sections:
        H(doc, t, 1); P(doc, b)
    out = BASE / "Sukuk-and-Green-Bond-Issuance-Guide.docx"
    doc.save(out); print(f"WROTE {out}")


def carve_out():
    doc = new_doc()
    cover(doc, "Carve-Out and Divestiture Playbook",
          "Partial IPO · Sale · Wind-Down · Asset Sale")
    sections = [
        ("1. Carve-Out IPO vs. Spinoff vs. Sale",
         "Carve-Out: Group retains majority post-IPO; raises primary "
         "capital, builds public-market currency, retains synergies. "
         "Spinoff: distribute to shareholders, tax-free if Section 355. "
         "Sale: cash exit, no continued exposure, certainty. Choice driven "
         "by valuation arbitrage, capital needs, brand strategy, and "
         "Group strategic intent."),
        ("2. Preparation",
         "Audited stand-alone financials (3 years), separate ERP run, "
         "stand-alone management team, customer and supplier "
         "communications plan, employee retention plan."),
        ("3. Stand-Alone P&L Construction",
         "Allocate corporate overhead conservatively to avoid post-deal "
         "earnings gap. Identify and quantify stranded costs that remain "
         "at Group. Negotiate TSA revenue back to Group to offset stranded "
         "costs."),
        ("4. Sale Process",
         "Auction vs. proprietary process. Use auction where multiple "
         "credible buyers exist; proprietary where price-sensitivity, "
         "regulatory risk, or strategic confidentiality argue for one-on-"
         "one. Banker-led process with full diligence room (VDR), "
         "management presentations, site visits."),
        ("5. Antitrust and Foreign Investment Review",
         "HSR (US), EU Merger Regulation, China SAMR, CFIUS for inbound "
         "US deals, FIRB (Australia), CMA (UK). Add 4-9 months to "
         "timeline for material reviews."),
        ("6. Brand Treatment Post-Sale",
         "Buyer typically retains the operating brand. AFANO endorsement "
         "footer is removed at closing. AFANO retains no licensing or "
         "brand interest absent specific transitional license."),
        ("7. Wind-Down",
         "Where sale is infeasible and continued operation is "
         "uneconomic, an orderly wind-down protects creditors, employees, "
         "and Group reputation. Use a specialist wind-down team; "
         "communicate openly with stakeholders."),
        ("8. Asset Sale vs. Equity Sale",
         "Equity sale is cleaner for Group (no historical liabilities), "
         "buyer prefers asset sale for step-up basis and liability "
         "control. Negotiate accordingly with appropriate tax leakage "
         "modeling."),
    ]
    for t, b in sections:
        H(doc, t, 1); P(doc, b)
    out = BASE / "Carve-Out-and-Divestiture-Playbook.docx"
    doc.save(out); print(f"WROTE {out}")


def main():
    BASE.mkdir(parents=True, exist_ok=True)
    ipo_readiness()
    spinoff()
    sukuk_green()
    carve_out()


if __name__ == "__main__":
    main()
