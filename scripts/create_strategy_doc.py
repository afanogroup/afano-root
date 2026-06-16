"""Generate AFANO-Group-Holdings-Strategy-v2.docx from canonical data."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import afano_data as A

OUT = Path(__file__).parent.parent / "00_Strategy" / "AFANO-Group-Holdings-Strategy-v2.docx"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
GOLD = RGBColor(0xC8, 0xA0, 0x4A)
GREY = RGBColor(0x55, 0x55, 0x55)


def set_style(doc):
    s = doc.styles["Normal"]
    s.font.name = "Calibri"
    s.font.size = Pt(11)


def H(doc, text, level=1, color=NAVY, size=None):
    p = doc.add_heading("", level=level)
    r = p.add_run(text)
    r.bold = True
    r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return p


def P(doc, text, bold=False, italic=False, color=None, size=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet" if level == 0 else "List Bullet 2")
    return p


def page_break(doc):
    doc.add_page_break()


def cover(doc):
    P(doc, "")
    P(doc, "")
    P(doc, "")
    P(doc, A.GROUP_NAME, bold=True, color=NAVY, size=36, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, A.GROUP_TAGLINE, italic=True, color=GREY, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, "")
    P(doc, "Strategy, Structure & Brand Architecture", color=GOLD, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, "")
    P(doc, "Version 2.0", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    P(doc, A.HOLDINGS_DOMAIN, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)


def executive_summary(doc):
    H(doc, "Executive Summary", level=1)
    P(doc,
      f"{A.GROUP_NAME} is structured as a multinational conglomerate in the "
      "tradition of Berkshire Hathaway, Mitsubishi Corporation, Samsung Group "
      "and Dangote Group. The Group operates a four-tier hierarchy:")
    bullet(doc, "Tier 1 — AFANO Group Holdings (the parent holding company).")
    bullet(doc, "Tier 2 — Strategic Business Groups (15 SBGs, each a wholly-owned intermediate holding entity).")
    bullet(doc, "Tier 3 — Business Units (separately incorporated operating subsidiaries).")
    bullet(doc, "Tier 4 — Services and Sub-Services delivered by each Business Unit.")
    P(doc,
      "This architecture is designed to support multiple subsidiaries, multiple "
      "brands, international operations, public-private partnerships, "
      "government contracts, foreign direct investment, mergers and "
      "acquisitions, and IPO-ready business units, while preserving brand "
      "independence so that no operating subsidiary depends on the AFANO name "
      "for its market reputation.")


def holdings_office(doc):
    H(doc, "Holdings Office — Centralised Functions", level=1)
    P(doc,
      "The Holdings Office is the corporate centre. It performs the functions "
      "common to all 15 SBGs and is the single interface between the Group and "
      "the capital markets, regulators, ratings agencies and sovereign "
      "investors.")
    for name, summary in A.HOLDINGS_OFFICE:
        H(doc, name, level=2, color=NAVY, size=14)
        P(doc, summary)


def governance(doc):
    H(doc, "Governance Principles", level=1)
    for name, body in A.GOVERNANCE_PRINCIPLES:
        H(doc, name, level=2, color=NAVY, size=14)
        P(doc, body)


def division(doc, d):
    num, name, color_hex, sub, glyph, summary, bus = d
    page_break(doc)
    color = RGBColor.from_string(color_hex.lstrip("#"))
    H(doc, f"SBG {num} — {name}", level=1, color=color, size=22)
    P(doc, summary, italic=True, color=GREY)
    P(doc, f"Independent brand domain: to be assigned per Brand Architecture (see Domain & Branding annex).")
    P(doc, f"Working subdomain for cross-reference: {sub}.{A.HOLDINGS_DOMAIN}", color=GREY, size=9)
    for bu_name, services in bus:
        H(doc, f"Business Unit — {bu_name}", level=2, color=color, size=14)
        for s in services:
            bullet(doc, s)


def live_properties(doc):
    H(doc, "Currently Operating Properties", level=1)
    P(doc, "These customer-facing properties are live or in progress and are "
           "retained under their existing brands; future Business Units adopt "
           "the Brand Independence Policy described above.")
    for label, url in A.LIVE_PROPERTIES:
        bullet(doc, f"{label}: {url}")


def expansion(doc):
    H(doc, "Annex A — Expansion Roadmap (Optional Future SBGs 16–19)", level=1)
    P(doc,
      "The following divisions are recommended as candidate elevations once "
      "their underlying Business Units reach scale. They are documented here "
      "so the Group's structure can absorb them without reorganisation.")
    for name, body in A.EXPANSION_ANNEX:
        H(doc, name, level=2, color=NAVY, size=14)
        P(doc, body)


def annex_b(doc):
    H(doc, "Annex B — Brand & Domain Independence", level=1)
    P(doc,
      "Each Business Unit is assigned an independent brand and domain that "
      "does not begin or end with 'afano' or 'cybit'. This permits the "
      "subsidiary to be spun off, sold, listed, or operated as a joint "
      "venture without rename risk to the customer-facing identity.")
    P(doc,
      "All operating subsidiary websites include a single 'An AFANO Group "
      "Company' endorsement footer with the AFANO Network Node mark, linking "
      "back to afanogroup.com. This is the entire scope of the AFANO "
      "endorsement; no further AFANO co-branding is permitted in customer-"
      "facing contexts.")
    P(doc,
      "The complete Domain Architecture, with 3–5 candidate names per SBG and "
      "per Business Unit and a RDAP availability check column, is delivered "
      "as a separate workbook in 01_Branding-Domains/Domain-Architecture.xlsx.")


def annex_c(doc):
    H(doc, "Annex C — Phased Build Roadmap", level=1)
    P(doc,
      "The Group's build sequence mirrors the proven phased approach used in "
      "regulated industries: Governance Foundation → Operations Launch → "
      "Growth & Expansion → Scale & IPO. Each phase has a folder of "
      "deliverables in this repository:")
    for label in [
        "03_Governance-Foundation/ — Holdings legal stack, board charters, capital budget, multi-jurisdiction registration.",
        "04_Operations-Launch/ — 15 capability statements, shared-services SOPs, IR playbook, ESG framework, cyber & data policy.",
        "05_Growth-Expansion/ — M&A playbook, international market entry, PPP playbook, sovereign procurement, FDI strategy.",
        "06_Scale-IPO/ — IPO readiness, subsidiary spinoff, sukuk & green bond issuance, carve-out & divestiture.",
    ]:
        bullet(doc, label)


def closing(doc):
    page_break(doc)
    H(doc, "Closing Statement", level=1)
    P(doc,
      "AFANO Group Holdings is a structure capable of supporting 50–100+ "
      "subsidiary companies, 200+ service lines, and expansion across "
      "multiple countries while maintaining clear governance, branding and "
      "operational accountability. Each Business Unit is designed to evolve "
      "into a standalone, IPO-ready company that can acquire subsidiaries "
      "and build its own market reputation independent of the AFANO name.")


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    set_style(doc)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
    cover(doc)
    executive_summary(doc)
    holdings_office(doc)
    governance(doc)
    live_properties(doc)
    for d in A.DIVISIONS:
        division(doc, d)
    page_break(doc)
    expansion(doc)
    annex_b(doc)
    annex_c(doc)
    closing(doc)
    doc.save(OUT)
    print(f"WROTE {OUT}")


if __name__ == "__main__":
    main()
